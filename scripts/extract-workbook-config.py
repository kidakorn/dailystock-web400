"""Reads settings that live inside Daily_Stock.xlsx and writes them as JSON the web app can use:
  data/process-seq.json    the hidden PROCESS_SEQ sheet (display order of processes and sub-processes)
  data/package-groups.json the IC2 / IC3 package name lists, taken from the Power Query filters IC2_TOTAL and IC3_TOTAL

Usage: python3 extract-workbook-config.py <Daily_Stock.xlsx> [output dir = ../data]
Needs: openpyxl. Rerun it if you change the sequence sheet or the package lists in Excel.
"""
import base64
import io
import json
import re
import struct
import sys
import zipfile

import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else 'Daily_Stock.xlsx'
OUT = sys.argv[2] if len(sys.argv) > 2 else '../data'


def power_query_code(path):
    """The Power Query (M) source is a base64 'DataMashup' blob stored in one of the customXml parts."""
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if not (name.startswith('customXml/item') and name.endswith('.xml')):
                continue
            raw = z.read(name)
            text = raw.decode('utf-16') if raw[:2] in (b'\xff\xfe', b'\xfe\xff') else raw.decode('utf-8-sig', 'ignore')
            hit = re.search(r'<DataMashup[^>]*>(.*?)</DataMashup>', text, re.S)
            if hit:
                data = base64.b64decode(hit.group(1).strip())
                (_, size) = struct.unpack('<II', data[:8])
                return zipfile.ZipFile(io.BytesIO(data[8:8 + size])).read('Formulas/Section1.m').decode('utf-8-sig')
    raise SystemExit('No Power Query code found in the workbook')


def package_names(m_code, query):
    block = [b for b in re.split(r'(?m)^(?=shared\s)', m_code) if b.startswith(f'shared {query} ')]
    if not block:
        raise SystemExit(f'Query {query} not found')
    return list(dict.fromkeys(re.findall(r'\[PACKACE\] = "([^"]+)"', block[0])))  # keeps order, drops repeats


m = power_query_code(SRC)
groups = {'IC2': package_names(m, 'IC2_TOTAL'), 'IC3': package_names(m, 'IC3_TOTAL')}

ws = openpyxl.load_workbook(SRC, data_only=True)['PROCESS_SEQ']
rows = ws[ws.tables['PROCESS_SEQ'].ref][1:]  # skip the header row
seq = []
for seq_no, label, proc, sub in ([c.value for c in r] for r in rows):
    if seq_no is None or proc is None:
        continue
    seq.append({'seq': int(seq_no), 'label': str(label or '').replace('\xa0', ' ').strip(),
                'processCode': int(proc), 'subProcessCode': '' if sub is None else str(int(sub))})

for name, data in [('process-seq.json', seq), ('package-groups.json', groups)]:
    with open(f'{OUT}/{name}', 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
print(f'process-seq: {len(seq)} entries | IC2: {len(groups["IC2"])} packages | IC3: {len(groups["IC3"])} packages')
